#!/usr/bin/env python3
"""Откуда взят принт: наши товары на Ozon ↔ товар-источник и магазин, у которого его скопировали.

Один постоянный CSV на тип товара — printflow/data/<товар>/reports/sources.csv; каждый запуск его
дообновляет (новых файлов не создаёт). В файл попадают только успешно выложенные товары: задача карточки
done и товар есть на Ozon (product_id).

    python -m printflow.scripts.source_report                      # все товары
    python -m printflow.scripts.source_report --product mousepad   # один
    python -m printflow.scripts.source_report --no-ozon            # без запросов в Ozon

Связка: external_id задачи = ref:<SKU источника> → по SKU в таблицах парсера (parser/prog/parser/*.xlsx)
название, ссылка, Seller ID, цена, рейтинг источника. Наш SKU и ссылка — из Ozon API (только чтение).

Быстро и без лишней работы: строки, уже записанные с нашим SKU, не пересчитываются; таблицы парсера
читаются, только если есть новые товары; в Ozon уходят только товары без SKU — пачками по 1000,
параллельно по всем кабинетам. CSV — UTF-8 с BOM и «;», открывается в Excel как есть.
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import json
import sqlite3
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

from printflow.config import PACKAGE_DIR, ROOT, load_app_settings, product_data_dir
from printflow.domain.models import CardSpec
from printflow.scripts._cli import Target, add_product_arg, target, utf8_console
from printflow.services.ozon.client import OzonClient, OzonError

PARSER_DIR = ROOT / "parser" / "prog" / "parser"
OZON_BATCH = 1000       # товаров в одном запросе /v3/product/info/list
OZON_PARALLEL = 4       # одновременных запросов в один кабинет
COLUMNS = [
    "Артикул товара", "Артикул дизайна", "Вариант", "Название модели (карточка)", "Ozon product_id", "Ozon SKU",
    "Наш товар на Ozon", "Принт (S3)", "Главное фото (S3)",
    "SKU источника", "Товар-источник", "Ссылка на источник", "Seller ID источника", "Магазин источника",
    "Цена источника", "Рейтинг", "Отзывы", "Фото источника", "Таблица парсера",
]
KEY, SKU_COL = "Артикул товара", "Ozon SKU"
SOURCE_FIELDS = ("Название", "Ссылка", "Seller ID", "Цена", "Рейтинг", "Отзывы", "Фото")


def report_path(product: str, data_dir: Path | None = None) -> Path:
    root = data_dir if data_dir is not None else load_app_settings().data_dir
    return product_data_dir(root, product) / "reports" / "sources.csv"


def read_report(path: Path) -> dict[str, dict[str, str]]:
    """Уже собранные строки: артикул товара → строка."""
    if not path.is_file():
        return {}
    with path.open(encoding="utf-8-sig", newline="") as f:
        return {row[KEY]: row for row in csv.DictReader(f, delimiter=";") if row.get(KEY)}


def write_report(path: Path, rows: list[dict[str, str]]) -> None:
    """Целиком и атомарно: сначала во временный файл, потом подмена — файл не бывает недописанным."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    with tmp.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    try:
        tmp.replace(path)
    except PermissionError:  # Windows: файл открыт в Excel и заблокирован
        tmp.unlink(missing_ok=True)
        raise SystemExit(f"Не удалось записать {path}: файл открыт в другой программе (Excel?). "
                         f"Закройте его и запустите отчёт снова.") from None


def read_sources(paths: list[Path]) -> dict[str, dict[str, str]]:
    """SKU → нужные поля строки парсера (+ имя файла). Более новые таблицы перекрывают старые."""
    sources: dict[str, dict[str, str]] = {}
    for path in sorted(paths, key=lambda p: p.stat().st_mtime):
        wb = load_workbook(path, read_only=True)
        rows = wb.active.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows)]
        if "SKU" in header:
            sku_i = header.index("SKU")
            idx = {f: header.index(f) for f in SOURCE_FIELDS if f in header}
            for row in rows:
                if row[sku_i]:
                    sources[str(row[sku_i]).strip()] = ({f: str(row[i] or "").strip() for f, i in idx.items()}
                                                        | {"_file": path.name})
        wb.close()
    return sources


def done_offers(db_path: Path, t: Target) -> list[dict[str, str]]:
    """Успешно выложенные товары (без данных источника): по строке на артикул товара."""
    db = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    rows = db.execute(
        "SELECT external_id, payload, result FROM jobs WHERE product = ? AND status = 'done'"
        " AND json_extract(payload, '$.kind') = 'card' ORDER BY number", (t.product.name,)).fetchall()
    db.close()
    out = []
    for external_id, payload_json, result_json in rows:
        res, payload = json.loads(result_json or "{}"), json.loads(payload_json)
        ozon = res.get("ozon") or {}
        offers, pids = ozon.get("offers") or [], ozon.get("product_ids") or []
        if ozon.get("dry_run") or not pids:
            continue
        variants = t.product.variants(CardSpec.from_payload(payload))
        images = res.get("images") or {}
        for i, offer in enumerate(offers):
            if i >= len(pids):
                break
            variant = variants[i] if i < len(variants) else ""
            photos = images.get(variant) or next(iter(images.values()), [])
            out.append({
                KEY: offer, "Артикул дизайна": res.get("article", ""), "Вариант": variant,
                "Название модели (карточка)": res.get("merge_group", ""), "Ozon product_id": str(pids[i]),
                "Принт (S3)": res.get("print", ""), "Главное фото (S3)": photos[0] if photos else "",
                "SKU источника": external_id[4:] if (external_id or "").startswith("ref:") else "",
                "_photo": payload.get("reference_url") or "",
            })
    return out


def add_source(row: dict[str, str], sources: dict[str, dict[str, str]]) -> None:
    src = sources.get(row["SKU источника"], {})
    seller = src.get("Seller ID", "")
    row.update({
        "Товар-источник": src.get("Название", ""), "Ссылка на источник": src.get("Ссылка", ""),
        "Seller ID источника": seller, "Магазин источника": f"https://www.ozon.ru/seller/{seller}/" if seller else "",
        "Цена источника": src.get("Цена", ""), "Рейтинг": src.get("Рейтинг", ""), "Отзывы": src.get("Отзывы", ""),
        "Фото источника": src.get("Фото", "") or row.get("_photo", ""), "Таблица парсера": src.get("_file", ""),
    })


async def ozon_skus(t: Target, product_ids: list[int]) -> dict[int, int]:
    """product_id → наш SKU: пачки по 1000, до OZON_PARALLEL одновременно."""
    client = OzonClient(t.settings.ozon)
    gate = asyncio.Semaphore(OZON_PARALLEL)

    async def batch(chunk: list[int]) -> dict[int, int]:
        async with gate:
            items = (await client.post("/v3/product/info/list", {"product_id": chunk})).get("items") or []
        return {int(i["id"]): int(i["sku"]) for i in items if i.get("sku")}

    found: dict[int, int] = {}
    for part in await asyncio.gather(*(batch(product_ids[s:s + OZON_BATCH])
                                       for s in range(0, len(product_ids), OZON_BATCH))):
        found.update(part)
    return found


async def all_skus(targets: dict[str, Target], need: dict[str, list[int]]) -> dict[str, dict[int, int]]:
    """SKU по всем кабинетам параллельно; кабинет без ключей или с ошибкой — пусто, с сообщением."""
    async def one(name: str) -> tuple[str, dict[int, int]]:
        t = targets[name]
        if not need[name]:
            return name, {}
        if not t.settings.ozon.configured:
            print(f"  ! {name}: нет ключей {t.product.env_prefix}_OZON_* — наши SKU не подтянуты")
            return name, {}
        try:
            return name, await ozon_skus(t, need[name])
        except OzonError as exc:
            print(f"  ! {name}: Ozon — {str(exc)[:200]}")
            return name, {}

    return dict(await asyncio.gather(*(one(n) for n in targets)))


def main() -> int:
    utf8_console()
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    add_product_arg(parser, required=False)
    parser.add_argument("--xlsx", action="append", type=Path,
                        help="таблица парсера (можно несколько); по умолчанию все *.xlsx из parser/prog/parser")
    parser.add_argument("--no-ozon", action="store_true", help="не запрашивать наши SKU в Ozon")
    args = parser.parse_args()
    started = time.monotonic()

    app = load_app_settings()
    names = [args.product] if args.product else list(app.products)
    targets = {name: target(name) for name in names}

    # строки: уже записанные (с SKU — как есть) + новые успешно выложенные
    reports, fresh = {}, {}
    for name in names:
        existing = read_report(report_path(name, app.data_dir))
        rows = []
        for row in done_offers(app.jobs_db, targets[name]):
            old = existing.get(row[KEY])
            rows.append(old if old and old.get(SKU_COL) else {**(old or {}), **row})
        reports[name] = rows
        fresh[name] = sum(1 for r in rows if r[KEY] not in existing)

    if any(not r.get("Товар-источник") and r.get("SKU источника") for rows in reports.values() for r in rows):
        sources = read_sources(args.xlsx or sorted(PARSER_DIR.glob("*.xlsx")))
        for rows in reports.values():
            for row in rows:
                if not row.get("Товар-источник") and row.get("SKU источника"):
                    add_source(row, sources)

    need = {n: sorted({int(r["Ozon product_id"]) for r in reports[n] if not r.get(SKU_COL)}) for n in names}
    skus = {n: {} for n in names} if args.no_ozon else asyncio.run(all_skus(targets, need))
    for name in names:
        for row in reports[name]:
            if not row.get(SKU_COL) and (sku := skus[name].get(int(row["Ozon product_id"]))):
                row[SKU_COL], row["Наш товар на Ozon"] = str(sku), f"https://www.ozon.ru/product/{sku}/"
        write_report(report_path(name, app.data_dir), reports[name])
        without_sku = sum(1 for r in reports[name] if not r.get(SKU_COL))
        print(f"{targets[name].product.title}: товаров {len(reports[name])} (новых {fresh[name]}), "
              f"без нашего SKU {without_sku} → {report_path(name, app.data_dir)}")
    print(f"Готово за {time.monotonic() - started:.1f} с")
    return 0


if __name__ == "__main__":
    sys.exit(main())
