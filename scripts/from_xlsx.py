#!/usr/bin/env python3
"""Поставить в очередь товары из XLSX парсера Ozon (parser/prog/parser) — рисунок копируется 1:1.

Для каждой строки главное фото товара (колонка «Фото») уходит в очередь как reference_url: GPT снимает
с фото принт (у шоппера — без сумки и фона, у коврика — всю поверхность), дальше галерея и карточка Ozon.

    python -m printflow.scripts.from_xlsx --product shopper  products_merged-3sellers_….xlsx --limit 10 --dry-run
    python -m printflow.scripts.from_xlsx --product mousepad products-mousepad_….xlsx --limit 9999
    python -m printflow.scripts.from_xlsx --product mousepad products-mousepad_….xlsx --offset 50 --limit 1   # 51-я строка

Строки уходят пачками по --batch-size (50) в POST /api/v1/<товар>/jobs/batch — одна транзакция на пачку.
Повторный запуск по тому же файлу дублей не создаёт (external_id = ref:<SKU>).
Нужен запущенный сервис (docker compose up -d).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import httpx
from openpyxl import load_workbook

from printflow.scripts._cli import add_product_arg, target, utf8_console


def read_rows(path: Path) -> list[dict[str, str]]:
    sheet = load_workbook(path, read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    return [{h: str(v or "").strip() for h, v in zip(header, row)} for row in rows]


def main() -> int:
    utf8_console()
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    add_product_arg(parser)
    parser.add_argument("xlsx", type=Path, help="XLSX из парсера Ozon")
    parser.add_argument("--limit", type=int, default=10, help="сколько товаров взять (по умолчанию 10)")
    parser.add_argument("--offset", type=int, default=0, help="сколько строк с фото пропустить сначала (по умолчанию 0)")
    parser.add_argument("--dry-run", action="store_true", help="сгенерировать, но не выкладывать на Ozon")
    parser.add_argument("--api", default=None, help="адрес API (по умолчанию SM_HOST:SM_PORT)")
    parser.add_argument("--name", default=None, help="название карточки на Ozon (по умолчанию — у товара)")
    parser.add_argument("--batch-size", type=int, default=50, help="строк в одном запросе (одна транзакция), до 500")
    parser.add_argument("--verbose", action="store_true", help="печатать каждую задачу, а не только итог пачки")
    args = parser.parse_args()

    t = target(args.product)
    token = t.settings.api.token
    if not token:
        sys.exit(f"Нет {t.product.env_prefix}_API_TOKEN в printflow/.env")
    name = args.name or t.product.default_card_name
    rows = [r for r in read_rows(args.xlsx) if r.get("Фото")][args.offset: args.offset + args.limit]
    if not rows:
        sys.exit(f"В {args.xlsx} нет строк с колонкой «Фото»")

    counts = {"+": 0, "=": 0, "!": 0}
    # trust_env=False: API локальный, системный прокси (HTTP_PROXY у VPN) ответил бы на него 503
    with httpx.Client(base_url=args.api or t.app.api_url, headers={"Authorization": f"Bearer {token}"},
                      timeout=t.settings.http.timeout_sec, trust_env=False) as client:
        for start in range(0, len(rows), args.batch_size):
            batch = rows[start:start + args.batch_size]
            body = {"jobs": [_card(row, args.dry_run, name) for row in batch]}
            try:
                response = client.post(f"{t.api_prefix}/jobs/batch", json=body)
                response.raise_for_status()
            except httpx.HTTPError as exc:
                counts["!"] += len(batch)
                detail = (f"HTTP {exc.response.status_code} {exc.response.text[:300]}"
                          if isinstance(exc, httpx.HTTPStatusError) else str(exc))
                print(f"  ! строки {start + 1}–{start + len(batch)} не поставлены: {type(exc).__name__} {detail}")
                continue
            for row, job in zip(batch, response.json()["jobs"]):
                mark = "+" if job["created"] else "="
                counts[mark] += 1
                if args.verbose:
                    print(f"  {mark} {job['id']}  {row.get('Название', '')[:70]}")
            print(f"[{start + len(batch)}/{len(rows)}] новых: {counts['+']}, уже были: {counts['=']}, ошибок: {counts['!']}")
    print(f"Поставлено новых задач: {counts['+']} из {len(rows)} (уже были в очереди: {counts['=']}, ошибок: {counts['!']})")
    return 0


def _card(row: dict[str, str], dry_run: bool, name: str) -> dict[str, object]:
    return {
        "reference_url": row["Фото"],
        "dry_run": dry_run,
        "external_id": f"ref:{row.get('SKU') or row['Фото']}",
        **({"product_name": name} if name else {}),
    }


if __name__ == "__main__":
    sys.exit(main())
